"""
Elenco dei livelli di Gotcha! in Excel: numero, immagine, nome, cosa fare, i 3 hint, Easter egg (nome e come
sbloccarlo). Dati da Tools/levels.json.

Immagini: se esiste la cartella Screenshots/ (Unity: menu Gotcha > Screenshot di tutti i livelli) usa i veri
screenshot level_001.png, ...; altrimenti la cartella indicata con --images (anteprime).

Uso:  python Tools/levels_xlsx.py [--images CARTELLA] [--out Gotcha_livelli.xlsx]
"""
import argparse, json, os
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EGG_REWARD = "Video (con audio) da condividere col link personale per scaricare il gioco; premio a scelta, una volta per segreto: 15 minuti di vite infinite oppure 5 hint."


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--images", default=None)
    ap.add_argument("--out", default=os.path.join(ROOT, "Gotcha_livelli.xlsx"))
    a = ap.parse_args()
    levels = json.load(open(os.path.join(ROOT, "Tools", "levels.json"), encoding="utf-8"))
    shots = os.path.join(ROOT, "Screenshots")
    real = os.path.isdir(shots) and any(f.startswith("level_") for f in os.listdir(shots))
    img_dir = shots if real else a.images

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2B2133")
    egg_fill = PatternFill("solid", fgColor="FFF1C2")
    thin = Side(style="thin", color="D0C8D8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    wrap = Alignment(wrap_text=True, vertical="center")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Livelli"
    cols = [("N°", 6), ("Screenshot" if real else "Anteprima", 18), ("Nome", 20), ("Cosa fare (soluzione)", 42),
            ("Hint 1", 36), ("Hint 2", 36), ("Hint 3 (soluzione)", 40), ("Easter egg", 10), ("Nome easter egg", 18), ("Come sbloccarlo", 42)]
    for i, (h, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = bold, head_fill, center, border
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 24
    tmp = os.path.join(ROOT, "Tools", ".xlsx_img")
    os.makedirs(tmp, exist_ok=True)
    for r, l in enumerate(levels, 2):
        hints = (l.get("hints") or [l["hint"]]) + ["", "", ""]
        vals = [l["n"], None, l["title"], l["desc"], hints[0], hints[1], hints[2],
                "Sì" if l["egg"] else "No", l["eggName"], l["eggHow"]]
        for i, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border = font, border
            c.alignment = center if i in (1, 8) else wrap
            if l["egg"] and i >= 8:
                c.fill = egg_fill
        ws.row_dimensions[r].height = 160
        if img_dir:
            for name in ("level_%03d.png" % l["n"], "%03d.png" % l["n"]):
                src = os.path.join(img_dir, name)
                if os.path.exists(src):
                    im = Image.open(src).convert("RGB")
                    im.thumbnail((116, 206))
                    dst = os.path.join(tmp, "%03d.png" % l["n"])
                    im.save(dst)
                    pic = XLImage(dst)
                    ws.add_image(pic, "B%d" % r)
                    break
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:J%d" % (len(levels) + 1)

    es = wb.create_sheet("Easter egg")
    ecols = [("Nome", 18), ("ID", 16), ("Livello", 26), ("Come sbloccarlo", 60), ("Condivisione e premio", 60)]
    for i, (h, w) in enumerate(ecols, 1):
        c = es.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = bold, head_fill, center, border
        es.column_dimensions[c.column_letter].width = w
    r = 2
    for l in levels:
        if not l["egg"]:
            continue
        for i, v in enumerate([l["eggName"], l["egg"], "%d - %s" % (l["n"], l["title"]), l["eggHow"], EGG_REWARD], 1):
            c = es.cell(row=r, column=i, value=v)
            c.font, c.border, c.alignment = font, border, wrap
        es.row_dimensions[r].height = 48
        r += 1
    es.freeze_panes = "A2"

    ns = wb.create_sheet("Note")
    notes = [
        "Gotcha! - elenco livelli (%d livelli)." % len(levels),
        ("Le immagini sono screenshot del gioco (Unity: Gotcha > Screenshot di tutti i livelli)." if real else
         "Le immagini sono ANTEPRIME composte con i disegni di ogni livello, non screenshot del gioco. Per i veri "
         "screenshot: in Unity menu Gotcha > Screenshot di tutti i livelli, poi rilanciare Tools/levels_xlsx.py."),
        "Hint: ogni livello ha 3 suggerimenti (spinta, più preciso, soluzione). Ogni suggerimento costa 1 hint; si parte con 5.",
        "Gli hint non si ricaricano: +1 guardando una pubblicità, 3 hint a 0,39 €, 15 hint a 1,49 €, oppure come premio per la condivisione.",
        "Easter egg: ogni segreto viene registrato (ultimi 14 secondi, con audio). Dopo la scoperta il gioco propone di condividere il video con un link personale per scaricare il gioco.",
        "Premio per la condivisione (una volta per ogni segreto): 15 minuti di vite infinite oppure 5 hint.",
        "Toni il carlino: un livello ogni 25 (25, 50, 75, 100, ...).",
    ]
    ns.column_dimensions["A"].width = 120
    for i, t in enumerate(notes, 1):
        c = ns.cell(row=i, column=1, value=t)
        c.font = Font(name="Arial", size=11, bold=(i == 1))
        c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(a.out)
    print("Salvato", a.out, "(screenshot)" if real else "(anteprime)" if img_dir else "(senza immagini)")


if __name__ == "__main__":
    main()
